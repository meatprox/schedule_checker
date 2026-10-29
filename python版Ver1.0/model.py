"""Excelの読み書きと、タスク・工数の計算。画面には依存しない。"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles.numbers import is_date_format
    from openpyxl.utils.datetime import from_excel
except ImportError as exc:
    raise SystemExit(
        "openpyxl が必要です。次を実行してください: py -3 -m pip install openpyxl"
    ) from exc

SLOT_COUNT = 10
MAX_SLOT_COUNT = 200
STYLES = ("solid", "dashed", "double", "triple", "quad", "quint", "none")
EFFORT_BY_STYLE = {"none": 0, "dashed": 0.5, "solid": 1, "double": 2, "triple": 3, "quad": 4, "quint": 5}
STYLE_FROM_LABEL = {
    "solid": "solid",
    "実線": "solid",
    "dashed": "dashed",
    "破線": "dashed",
    "double": "double",
    "二重線": "double",
    "triple": "triple",
    "三重線": "triple",
    "quad": "quad",
    "四重線": "quad",
    "quint": "quint",
    "五重線": "quint",
    "none": "none",
    "線無し": "none",
    "線なし": "none",
    "無し": "none",
    "なし": "none",
}
NAME_LABELS = {"タスク名", "タスク", "予定名", "項目名"}
DATE_LABELS = {"日付", "日時", "年月日"}
TIME_LABELS = {"時間", "時刻", "開始時刻"}
DESC_LABELS = {"説明", "内容", "メモ", "備考", "詳細"}
SLOT_LABELS = {"スロット", "レーン", "段"}
LINE_LABELS = {"線", "線種", "線の状態", "工数"}
COLOR_LABELS = {"色", "カラー", "color"}
LINE_COLORS = (
    "#0e6b67",
    "#c2410c",
    "#1d4ed8",
    "#7c3aed",
    "#be123c",
    "#15803d",
    "#b45309",
    "#0369a1",
    "#a21caf",
    "#4d7c0f",
)
ERA_BASE = {"令和": 2018, "平成": 1988, "昭和": 1925}

SAMPLE_POINTS = (
    ("受付", 2026, 10, 21, 8, 30, "来客対応"),
    ("朝会", 2026, 10, 21, 9, 0, "本日の目標を共有"),
    ("朝会", 2026, 10, 21, 9, 15, "課題の洗い出し"),
    ("朝会", 2026, 10, 21, 9, 40, "終了"),
    ("設計レビュー", 2026, 10, 21, 10, 0, "資料の説明"),
    ("設計レビュー", 2026, 10, 21, 11, 10, "指摘の確認"),
    ("設計レビュー", 2026, 10, 21, 12, 0, "レビューを閉じる"),
    ("移動", 2026, 10, 21, 10, 30, "会議室へ移動"),
    ("移動", 2026, 10, 21, 11, 0, "到着"),
    ("昼休憩", 2026, 10, 21, 12, 10, "外出"),
    ("昼休憩", 2026, 10, 21, 13, 0, "帰席"),
    ("実装", 2026, 10, 21, 13, 10, "着手"),
    ("実装", 2026, 10, 21, 15, 0, "中間確認"),
    ("実装", 2026, 10, 21, 17, 40, "一区切り"),
    ("定例", 2026, 10, 21, 15, 0, "進捗の共有"),
    ("定例", 2026, 10, 21, 16, 20, "決定事項の確認"),
    ("相談", 2026, 10, 21, 15, 30, "仕様の確認"),
    ("相談", 2026, 10, 21, 16, 0, "対応方針を決める"),
    ("報告", 2026, 10, 21, 17, 50, "文面を作成"),
    ("報告", 2026, 10, 21, 18, 20, "送付"),
    ("振り返り", 2026, 10, 22, 9, 0, "昨日の確認"),
    ("振り返り", 2026, 10, 22, 9, 40, "次の予定を整理"),
)


@dataclass
class Parts:
    y: int = 0
    m: int = 0
    d: int = 0
    H: int = 0
    M: int = 0
    S: int = 0
    has_date: bool = False
    has_time: bool = False


@dataclass
class Point:
    id: str
    time: datetime
    description: str
    row: int
    outgoing_style: str = ""


@dataclass
class Segment:
    style: str = "solid"


@dataclass
class Task:
    id: str
    name: str
    slot: int = 0
    slot_assigned: bool = False
    points: list[Point] = field(default_factory=list)
    segments: list[Segment] = field(default_factory=list)
    color: str = ""


@dataclass
class ParseResult:
    tasks: list[Task]
    skipped: int = 0
    sheet_name: str = ""
    slot_names: list[str] = field(default_factory=list)


@dataclass
class EffortSpan:
    start: datetime
    end: datetime
    total: float


def clamp(low, value, high):
    return min(high, max(low, value))


def clean_text(value) -> str:
    text = str(value if value is not None else "")
    text = text.replace("\u3000", " ").replace("\r\n", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    return text.strip()


def normalize_width(value) -> str:
    text = clean_text(value)
    text = text.translate(str.maketrans("０１２３４５６７８９：／．－", "0123456789:/-."))
    return re.sub(r"[ \t]+", " ", text)


def header_token(text: str) -> str:
    token = normalize_width(text)
    token = re.sub(r"\s", "", token)
    return re.sub(r"[：:*＊]+$", "", token)


def label_kind(text: str) -> str:
    token = header_token(text)
    if not token:
        return ""
    if token in NAME_LABELS:
        return "name"
    if token in DATE_LABELS:
        return "date"
    if token in TIME_LABELS:
        return "time"
    if token in DESC_LABELS:
        return "desc"
    if token in SLOT_LABELS:
        return "slot"
    if token in LINE_LABELS:
        return "line"
    if token in COLOR_LABELS:
        return "color"
    return ""


def style_from_effort(value) -> str:
    try:
        effort = round(float(value) * 10) / 10
    except (TypeError, ValueError):
        return ""
    return {0: "none", 0.5: "dashed", 1: "solid", 2: "double", 3: "triple", 4: "quad", 5: "quint"}.get(effort, "")


def parse_line_value(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return style_from_effort(value)
    token = header_token(normalize_width(value))
    if token in STYLE_FROM_LABEL:
        return STYLE_FROM_LABEL[token]
    return style_from_effort(token)


def parse_color_value(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        token = header_token(value).lower()
        if token in LINE_COLORS:
            return token
    else:
        token = value
    try:
        number = int(round(float(token)))
    except (TypeError, ValueError):
        return ""
    if 1 <= number <= len(LINE_COLORS):
        return LINE_COLORS[number - 1]
    return ""


def parse_slot_value(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        raw = float(value)
    else:
        token = header_token(normalize_width(value))
        try:
            raw = float(token)
        except ValueError:
            return None
    slot = int(round(raw))
    if slot < 1 or slot > MAX_SLOT_COUNT:
        return None
    return slot - 1


def slot_count_for(tasks, minimum: int = SLOT_COUNT) -> int:
    count = minimum
    for task in tasks:
        if task.slot_assigned and isinstance(task.slot, int):
            count = max(count, task.slot + 1)
    return count


def normalize_clock(hours: int, minutes: int, seconds: int) -> tuple[int, int, int]:
    seconds = int(round(seconds))
    if seconds >= 60:
        seconds = 0
        minutes += 1
    if minutes >= 60:
        minutes = 0
        hours += 1
    if hours >= 24:
        return 23, 59, 59
    return hours, minutes, seconds


def apply_meridiem(hours: int, marker: str) -> int:
    token = (marker or "").upper()
    if token in {"PM", "午後"} and hours < 12:
        return hours + 12
    if token in {"AM", "午前"} and hours == 12:
        return 0
    return hours


def clock_parts(hours: int, minutes: int, seconds: int) -> Parts:
    hours, minutes, seconds = normalize_clock(hours, minutes, seconds)
    return Parts(H=hours, M=minutes, S=seconds, has_time=True)


def parse_clock(text: str):
    matched = re.match(
        r"^(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?(?:\s*(AM|PM|午前|午後))?$",
        text,
        flags=re.IGNORECASE,
    )
    if matched:
        return clock_parts(
            apply_meridiem(int(matched.group(1)), matched.group(4) or ""),
            int(matched.group(2)),
            int(matched.group(3) or 0),
        )
    dot = re.match(r"^(\d{1,2})\.(\d{2})$", text)
    if dot:
        return clock_parts(int(dot.group(1)), int(dot.group(2)), 0)
    jp_time = re.match(
        r"^(午前|午後|AM|PM)?\s*(\d{1,2})時(?:(\d{1,2})分)?(?:(\d{1,2})秒)?$",
        text,
        flags=re.IGNORECASE,
    )
    if jp_time:
        return clock_parts(
            apply_meridiem(int(jp_time.group(2)), jp_time.group(1) or ""),
            int(jp_time.group(3) or 0),
            int(jp_time.group(4) or 0),
        )
    return None


def parse_loose_text(raw: str):
    text = normalize_width(raw)
    text = re.sub(r"[（(][日月火水木金土祝][)）]$", "", text)
    if not text or label_kind(text):
        return None

    era = re.match(r"^(令和|平成|昭和)(\d{1,2}|元)年(\d{1,2})月(\d{1,2})日(?:\s*(.+))?$", text)
    if era:
        year_number = 1 if era.group(2) == "元" else int(era.group(2))
        tail = parse_loose_text(era.group(5)) if era.group(5) else None
        return Parts(
            y=ERA_BASE[era.group(1)] + year_number,
            m=int(era.group(3)),
            d=int(era.group(4)),
            H=tail.H if tail else 0,
            M=tail.M if tail else 0,
            S=tail.S if tail else 0,
            has_date=True,
            has_time=bool(tail and tail.has_time),
        )

    date_time = re.match(
        r"^(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})(?:[ T](\d{1,2}):(\d{1,2})(?::(\d{1,2}))?)?$",
        text,
    )
    if date_time:
        has_time = date_time.group(4) is not None
        return Parts(
            y=int(date_time.group(1)),
            m=int(date_time.group(2)),
            d=int(date_time.group(3)),
            H=int(date_time.group(4)) if has_time else 0,
            M=int(date_time.group(5)) if has_time else 0,
            S=int(date_time.group(6) or 0) if has_time else 0,
            has_date=True,
            has_time=has_time,
        )

    jp = re.match(
        r"^(\d{4})年(\d{1,2})月(\d{1,2})日(?:\s*(午前|午後)?\s*(\d{1,2})時(?:(\d{1,2})分)?(?:(\d{1,2})秒)?)?$",
        text,
    )
    if jp:
        has_time = jp.group(5) is not None
        return Parts(
            y=int(jp.group(1)),
            m=int(jp.group(2)),
            d=int(jp.group(3)),
            H=apply_meridiem(int(jp.group(5)), jp.group(4) or "") if has_time else 0,
            M=int(jp.group(6) or 0) if has_time else 0,
            S=int(jp.group(7) or 0) if has_time else 0,
            has_date=True,
            has_time=has_time,
        )

    month_day = re.match(r"^(\d{1,2})[/\-.](\d{1,2})$", text)
    if month_day:
        return Parts(m=int(month_day.group(1)), d=int(month_day.group(2)), has_date=True)
    jp_month_day = re.match(r"^(\d{1,2})月(\d{1,2})日$", text)
    if jp_month_day:
        return Parts(m=int(jp_month_day.group(1)), d=int(jp_month_day.group(2)), has_date=True)
    return parse_clock(text)


def valid_date_parts(parts: Parts | None) -> bool:
    if not parts or not parts.has_date:
        return False
    if parts.m < 1 or parts.m > 12 or parts.d < 1 or parts.d > 31:
        return False
    if parts.y and (parts.y < 1901 or parts.y > 2100):
        return False
    return True


def valid_time_parts(parts: Parts | None) -> bool:
    if not parts or not parts.has_time:
        return False
    return 0 <= parts.H <= 23 and 0 <= parts.M <= 59 and 0 <= parts.S <= 59


def make_date(year: int, month: int, day: int, hours: int, minutes: int, seconds: int):
    try:
        value = datetime(year, month, day, hours, minutes, seconds)
    except ValueError:
        return None
    if (value.year, value.month, value.day) != (year, month, day):
        return None
    return value


def using_1904(workbook: Workbook) -> bool:
    epoch = getattr(workbook, "epoch", None)
    return bool(epoch and getattr(epoch, "year", None) == 1904)


def analyze_value(value, number_format: str, date1904: bool):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        hours, minutes, seconds = normalize_clock(value.hour, value.minute, value.second)
        return Parts(
            y=value.year,
            m=value.month,
            d=value.day,
            H=hours,
            M=minutes,
            S=seconds,
            has_date=True,
            has_time=True,
        )
    if isinstance(value, date):
        return Parts(y=value.year, m=value.month, d=value.day, has_date=True)
    if isinstance(value, time):
        return clock_parts(value.hour, value.minute, value.second)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = float(value)
        format_is_date = bool(number_format and is_date_format(number_format))
        looks_like_date = format_is_date or (0 <= serial < 1) or (20000 <= serial < 80000)
        if not looks_like_date:
            return None
        try:
            converted = from_excel(serial, epoch=datetime(1904, 1, 1) if date1904 else datetime(1899, 12, 30))
        except (OverflowError, ValueError, TypeError):
            return None
        if isinstance(converted, datetime):
            hours, minutes, seconds = normalize_clock(converted.hour, converted.minute, converted.second)
            fraction = abs(serial - int(serial))
            return Parts(
                y=converted.year,
                m=converted.month,
                d=converted.day,
                H=hours,
                M=minutes,
                S=seconds,
                has_date=serial >= 1,
                has_time=fraction > 1e-8 or (0 <= serial < 1),
            )
        if isinstance(converted, time):
            return clock_parts(converted.hour, converted.minute, converted.second)
        return None
    if isinstance(value, str):
        return parse_loose_text(value)
    return None


def analyze_cell(cell, date1904: bool):
    if cell is None:
        return None
    return analyze_value(cell.value, cell.number_format or "", date1904)


def cell_text(cell) -> str:
    if cell is None or cell.value is None:
        return ""
    value = cell.value
    if isinstance(value, str):
        return normalize_width(value)
    if isinstance(value, datetime):
        return normalize_width(value.strftime("%Y/%m/%d %H:%M:%S"))
    if isinstance(value, date):
        return normalize_width(value.strftime("%Y/%m/%d"))
    if isinstance(value, time):
        return normalize_width(value.strftime("%H:%M:%S"))
    if isinstance(value, (int, float, bool)):
        return normalize_width(value)
    return ""


def find_layout(sheet):
    last_row = min(sheet.max_row or 1, 30)
    last_col = min(sheet.max_column or 1, 12)
    for row in range(1, last_row + 1):
        for col in range(1, last_col + 1):
            if label_kind(cell_text(sheet.cell(row, col))) != "name":
                continue
            layout = {
                "name_col": col,
                "date_col": col + 1,
                "time_col": col + 2,
                "desc_col": col + 3,
                "slot_col": None,
                "line_col": None,
                "color_col": None,
                "data_row": row + 1,
                "labeled": False,
            }
            for probe in range(col + 1, min(sheet.max_column or col, col + 8) + 1):
                kind = label_kind(cell_text(sheet.cell(row, probe)))
                if kind == "date":
                    layout["date_col"] = probe
                elif kind == "time":
                    layout["time_col"] = probe
                elif kind == "desc":
                    layout["desc_col"] = probe
                elif kind == "slot":
                    layout["slot_col"] = probe
                elif kind == "line":
                    layout["line_col"] = probe
                elif kind == "color":
                    layout["color_col"] = probe
                if kind:
                    layout["labeled"] = True
            if layout["slot_col"] is None:
                candidate = col + 4
                used = {layout["date_col"], layout["time_col"], layout["desc_col"], layout["line_col"]}
                if candidate not in used and not cell_text(sheet.cell(row, candidate)):
                    layout["slot_col"] = candidate
            if layout["labeled"]:
                return layout
    return {
        "name_col": 1,
        "date_col": 2,
        "time_col": 3,
        "desc_col": 4,
        "slot_col": 5,
        "line_col": None,
        "color_col": None,
        "data_row": 1,
        "labeled": False,
    }


def parse_sheet(sheet, date1904: bool):
    if not sheet.max_row:
        return [], 0
    layout = find_layout(sheet)
    last_row = min(sheet.max_row, layout["data_row"] + 4999)
    pending = []
    carried_name = ""
    carried_date = None
    group_break = True
    skipped = 0

    for row in range(layout["data_row"], last_row + 1):
        name_cell = sheet.cell(row, layout["name_col"])
        date_cell = sheet.cell(row, layout["date_col"])
        time_cell = sheet.cell(row, layout["time_col"])
        desc_cell = sheet.cell(row, layout["desc_col"])
        slot_value = None
        if layout["slot_col"] is not None:
            slot_value = parse_slot_value(sheet.cell(row, layout["slot_col"]).value)
        line_value = ""
        if layout["line_col"] is not None:
            line_value = parse_line_value(sheet.cell(row, layout["line_col"]).value)
        color_value = ""
        if layout["color_col"] is not None:
            color_value = parse_color_value(sheet.cell(row, layout["color_col"]).value)
        name_text = cell_text(name_cell)
        desc_text = cell_text(desc_cell)
        date_info = analyze_cell(date_cell, date1904)
        time_info = analyze_cell(time_cell, date1904)
        fully_empty = (
            not name_text
            and not desc_text
            and not date_info
            and not time_info
            and slot_value is None
            and not line_value
        )
        if fully_empty or label_kind(name_text) == "name":
            if fully_empty:
                carried_name = ""
                group_break = True
            continue
        if name_text:
            if name_text != carried_name:
                group_break = True
            carried_name = name_text
        if valid_date_parts(date_info):
            carried_date = (date_info.y, date_info.m, date_info.d)
        elif valid_date_parts(time_info):
            carried_date = (time_info.y, time_info.m, time_info.d)
        clock = time_info if valid_time_parts(time_info) else date_info if valid_time_parts(date_info) else None
        if not carried_name or clock is None or carried_date is None:
            if name_text or desc_text or date_info or time_info:
                skipped += 1
            continue
        pending.append(
            {
                "name": carried_name,
                "date": carried_date,
                "time": (clock.H, clock.M, clock.S),
                "description": "" if label_kind(desc_text) else desc_text,
                "row": row,
                "break_before": group_break,
                "slot": slot_value,
                "line_style": line_value,
                "color": color_value,
            }
        )
        group_break = False

    known_year = next((item["date"][0] for item in pending if 1901 <= item["date"][0] <= 2100), datetime.now().year)
    tasks: list[Task] = []
    for item in pending:
        year = item["date"][0] if item["date"][0] >= 1901 else known_year
        stamp = make_date(year, item["date"][1], item["date"][2], *item["time"])
        if stamp is None:
            skipped += 1
            continue
        previous = tasks[-1] if tasks else None
        if item["break_before"] or previous is None or previous.name != item["name"]:
            tasks.append(Task(id=f"t{len(tasks) + 1}", name=item["name"]))
        task = tasks[-1]
        if item["slot"] is not None and not task.slot_assigned:
            task.slot = item["slot"]
            task.slot_assigned = True
        if item["color"] and not task.color:
            task.color = item["color"]
        task.points.append(
            Point(
                id=f"p{item['row']}",
                time=stamp,
                description=item["description"],
                row=item["row"],
                outgoing_style=item["line_style"],
            )
        )

    for task in tasks:
        task.points.sort(key=lambda point: (point.time, point.row))
        task.segments = [
            Segment(style=point.outgoing_style if point.outgoing_style in STYLES else "solid")
            for point in task.points[:-1]
        ]
        for point in task.points:
            point.outgoing_style = ""
    return [task for task in tasks if task.points], skipped


def parse_workbook(source) -> ParseResult:
    opened = None
    if isinstance(source, Workbook):
        workbook = source
    else:
        workbook = load_workbook(source, data_only=False, rich_text=False)
        opened = workbook
    try:
        date1904 = using_1904(workbook)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            if not sheet.max_row:
                continue
            tasks, skipped = parse_sheet(sheet, date1904)
            if tasks:
                return ParseResult(
                    tasks=tasks,
                    skipped=skipped,
                    sheet_name=sheet_name,
                    slot_names=read_slot_names(sheet),
                )
        names = workbook.sheetnames
        return ParseResult(tasks=[], skipped=0, sheet_name=names[0] if names else "")
    finally:
        if opened is not None:
            opened.close()


def task_span(task: Task):
    start = task.points[0].time
    last = task.points[-1].time
    minimum = start + timedelta(minutes=1)
    return start, last if last > minimum else minimum


def ranges_overlap(start, end, other_start, other_end) -> bool:
    return start < other_end and end > other_start


def assign_slots(tasks: list[Task], slot_count: int = SLOT_COUNT) -> list[Task]:
    spans: list[list[tuple[datetime, datetime]]] = [[] for _ in range(slot_count)]
    for task in tasks:
        if not task.slot_assigned or not task.points:
            continue
        task.slot = clamp(0, task.slot, slot_count - 1)
        spans[task.slot].append(task_span(task))

    ordered = sorted(
        (task for task in tasks if not task.slot_assigned and task.points),
        key=lambda task: (task.points[0].time, task.id),
    )
    for task in ordered:
        start, end = task_span(task)
        chosen = next(
            (index for index, listed in enumerate(spans) if not any(ranges_overlap(start, end, begin, finish) for begin, finish in listed)),
            -1,
        )
        if chosen < 0:
            ranked = sorted(
                (
                    (
                        sum(1 for begin, finish in listed if ranges_overlap(start, end, begin, finish)),
                        index,
                    )
                    for index, listed in enumerate(spans)
                )
            )
            chosen = ranked[0][1]
        task.slot = chosen
        spans[chosen].append((start, end))
    return tasks


def safe_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\\/?*\[\]:]", " ", str(name or "予定")).strip()[:31]
    return cleaned or "予定"


def export_file_name(file_name: str) -> str:
    base = re.sub(r"\.(xlsx|xlsm|xls)$", "", str(file_name or "スケジュール"), flags=re.IGNORECASE)
    return f"{base}_編集.xlsx"


def read_slot_names(sheet) -> list[str]:
    found = None
    last_row = min(sheet.max_row or 1, 30)
    last_col = min(sheet.max_column or 1, 12)
    for row in range(1, last_row + 1):
        for col in range(1, last_col + 1):
            if header_token(cell_text(sheet.cell(row, col))) == "スロット名":
                found = (row, col)
                break
        if found:
            break
    if found is None:
        return []
    header_row, col = found
    names = []
    for row in range(header_row + 1, (sheet.max_row or header_row) + 1):
        number = sheet.cell(row, col).value
        if number is None or number == "":
            break
        try:
            slot_number = int(round(float(number)))
        except (TypeError, ValueError):
            break
        if slot_number != len(names) + 1:
            break
        names.append(cell_text(sheet.cell(row, col + 1)))
    return names


def tasks_to_workbook(tasks: list[Task], sheet_name: str = "予定", slot_names=None) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = safe_sheet_name(sheet_name)
    headers = ("タスク名", "日付", "時間", "説明", "スロット", "工数")
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column, header)
    row_index = 2
    # タスクは画面上の並びのまま書く。時刻順に並べ替えると、読み込み後の段と重なり順が変わる。
    for task_index, task in enumerate(tasks):
        if not task.points:
            continue
        if row_index > 2 and task.name == tasks[task_index - 1].name:
            row_index += 1
        for index, point in enumerate(task.points):
            sheet.cell(row_index, 1, task.name)
            date_cell = sheet.cell(row_index, 2, point.time.date())
            date_cell.number_format = "yyyy/mm/dd"
            time_cell = sheet.cell(row_index, 3, point.time.time().replace(microsecond=0))
            time_cell.number_format = "hh:mm"
            sheet.cell(row_index, 4, point.description or "")
            sheet.cell(row_index, 5, (task.slot if isinstance(task.slot, int) else 0) + 1)
            if index < len(task.points) - 1:
                style = task.segments[index].style if index < len(task.segments) else "solid"
                sheet.cell(row_index, 6, EFFORT_BY_STYLE.get(style, 1))
            if task.color in LINE_COLORS:
                sheet.cell(1, 7, "色")
                sheet.cell(row_index, 7, LINE_COLORS.index(task.color) + 1)
            row_index += 1
    if slot_names:
        sheet.cell(1, 8, "スロット名")
        for index, name in enumerate(slot_names):
            sheet.cell(index + 2, 8, index + 1)
            text = str(name or "").strip()
            if text and text != str(index + 1):
                sheet.cell(index + 2, 9, text)
    widths = (16, 14, 10, 24, 10, 10)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    return workbook


def save_tasks(tasks: list[Task], sheet_name: str, path, slot_names=None) -> None:
    workbook = tasks_to_workbook(tasks, sheet_name, slot_names)
    workbook.save(path)
    workbook.close()


def effort_spans(tasks: list[Task]) -> list[EffortSpan]:
    events = []
    for task in tasks:
        for index, point in enumerate(task.points[:-1]):
            end = task.points[index + 1].time
            if end <= point.time:
                continue
            style = task.segments[index].style if index < len(task.segments) else "solid"
            if style not in EFFORT_BY_STYLE:
                style = "solid"
            effort = EFFORT_BY_STYLE[style]
            if effort == 0:
                continue
            events.append((point.time, effort))
            events.append((end, -effort))
    events.sort(key=lambda item: (item[0], item[1]))
    spans = []
    total = 0.0
    previous = None
    for moment, delta in events:
        if previous is not None and moment > previous and total > 0:
            spans.append(EffortSpan(start=previous, end=moment, total=total))
        total = round((total + delta) * 10) / 10
        previous = moment
    return spans


def clone_tasks(tasks: list[Task]) -> list[Task]:
    return [
        Task(
            id=task.id,
            name=task.name,
            slot=task.slot,
            slot_assigned=task.slot_assigned,
            points=[
                Point(id=point.id, time=point.time, description=point.description, row=point.row)
                for point in task.points
            ],
            segments=[Segment(style=segment.style) for segment in task.segments],
            color=task.color,
        )
        for task in tasks
    ]


def same_tasks(current: list[Task], snapshot: list[Task]) -> bool:
    if len(current) != len(snapshot):
        return False
    for task, other in zip(current, snapshot):
        if task.id != other.id or task.slot != other.slot or task.color != other.color:
            return False
        if len(task.segments) != len(other.segments):
            return False
        if any(segment.style != other.segments[index].style for index, segment in enumerate(task.segments)):
            return False
        if any(point.time != other.points[index].time for index, point in enumerate(task.points)):
            return False
    return True


def sample_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "予定"
    for column, header in enumerate(("タスク名", "日付", "時間", "説明"), start=1):
        sheet.cell(1, column, header)
    for index, point in enumerate(SAMPLE_POINTS, start=2):
        name, year, month, day, hour, minute, description = point
        sheet.cell(index, 1, name)
        date_cell = sheet.cell(index, 2, date(year, month, day))
        date_cell.number_format = "yyyy/mm/dd"
        time_cell = sheet.cell(index, 3, time(hour, minute))
        time_cell.number_format = "hh:mm"
        sheet.cell(index, 4, description)
    return workbook


def bundled_sample_path() -> Path:
    return Path(__file__).resolve().parent.parent / "public" / "sample.xlsx"

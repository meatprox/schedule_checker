"""今表示している予定から、ガントチャートの Excel を作る。予定データ自体は変えない。"""

from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from model import EFFORT_BY_STYLE, effort_spans

MAX_CELLS = 4000


class GanttRangeError(ValueError):
    """目盛が細かすぎて、1枚の表に収まらない。"""


def grid_origin(tasks) -> datetime | None:
    moments = [point.time for task in tasks for point in task.points]
    if not moments:
        return None
    first = min(moments)
    return datetime(first.year, first.month, first.day)


def step_size(amount: int, unit: str) -> timedelta:
    if unit == "day":
        return timedelta(days=amount)
    return timedelta(hours=amount)


def cell_index(moment: datetime, origin: datetime, step: timedelta) -> int:
    return int((moment - origin).total_seconds() // step.total_seconds())


def cell_start(index: int, origin: datetime, step: timedelta) -> datetime:
    return origin + step * index


def snap_time(moment: datetime, origin: datetime, step: timedelta, align: str) -> datetime:
    """点を、セルの先頭か末尾の時刻へ寄せる。自由は前半なら先頭、後半なら末尾。"""
    index = cell_index(moment, origin, step)
    start = cell_start(index, origin, step)
    end = start + step
    if align == "start":
        return start
    if align == "end":
        return end
    midpoint = start + step / 2
    if moment < midpoint:
        return start
    return end


def resolve_aligns(align: str | None, start_align: str | None, middle_align: str | None, end_align: str | None) -> tuple[str, str, str]:
    base = "free" if align in {None, "auto"} else align
    chosen = []
    for value in (start_align, middle_align, end_align):
        if value in {None, "auto"}:
            chosen.append(base if value is None else "free")
        else:
            chosen.append(value)
    return chosen[0], chosen[1], chosen[2]


def shade_alpha(effort: float) -> float:
    """工数5を線の色そのものにし、小さい工数ほど白に近づける。"""
    capped = min(max(effort, 0.0), 5.0)
    if capped <= 0:
        return 0.0
    if capped < 0.5:
        return 0.40 * (capped / 0.5)
    stops = ((0.5, 0.40), (1.0, 0.55), (2.0, 0.68), (3.0, 0.80), (4.0, 0.90), (5.0, 1.0))
    for index in range(1, len(stops)):
        left_effort, left_alpha = stops[index - 1]
        right_effort, right_alpha = stops[index]
        if capped <= right_effort:
            weight = (capped - left_effort) / (right_effort - left_effort)
            return left_alpha + (right_alpha - left_alpha) * weight
    return 1.0


def shade_color(color: str, effort: float) -> str:
    text = color.lstrip("#")
    channels = [int(text[index : index + 2], 16) for index in (0, 2, 4)]
    alpha = shade_alpha(effort)
    mixed = [round(255 + (channel - 255) * alpha) for channel in channels]
    return "{:02X}{:02X}{:02X}".format(*mixed)


def painted_cells(
    task,
    origin: datetime,
    step: timedelta,
    align: str | None,
    color: str,
    shade: bool = False,
    *,
    start_align: str | None = None,
    middle_align: str | None = None,
    end_align: str | None = None,
) -> dict[int, str]:
    """線の各区間を、寄せたあと半開区間 [左, 右) のセルに塗る。工数0は塗らない。"""
    points = task.points
    if not points:
        return {}
    start_align, middle_align, end_align = resolve_aligns(align, start_align, middle_align, end_align)
    if len(points) == 1:
        return _paint_single_point(points[0].time, origin, step, color, start_align, end_align)
    roles = [start_align if index == 0 else end_align if index == len(points) - 1 else middle_align for index in range(len(points))]
    snaps = [snap_time(point.time, origin, step, role) for point, role in zip(points, roles)]
    fills: dict[int, str] = {}
    for index, segment in enumerate(task.segments):
        if index + 1 >= len(points):
            continue
        effort = EFFORT_BY_STYLE.get(segment.style, 1)
        if effort <= 0:
            continue
        left = snaps[index]
        right = snaps[index + 1]
        if right <= left:
            continue
        chosen = shade_color(color, effort) if shade else color
        first = cell_index(left, origin, step)
        last = cell_index(right - timedelta(microseconds=1), origin, step)
        for cell in range(first, last + 1):
            fills[cell] = chosen
    return fills


def _paint_single_point(moment: datetime, origin: datetime, step: timedelta, color: str, start_align: str, end_align: str) -> dict[int, str]:
    """点が1つの線は、始点が前寄せで終点が後寄せのときだけ、そのセルを塗る。"""
    if start_align != "start" or end_align != "end":
        return {}
    return {cell_index(moment, origin, step): color}


def cell_effort(spans, start: datetime, end: datetime) -> float | None:
    overlapped = [span.total for span in spans if span.start < end and span.end > start]
    if not overlapped:
        return None
    return max(overlapped)


def slot_label(slot: int, names) -> str:
    """画面左の欄と同じ。書き換えた文字があればそれ、なければ番号。"""
    if names and 0 <= slot < len(names) and str(names[slot]).strip():
        return str(names[slot]).strip()
    return str(slot + 1)


def gantt_file_name(file_name: str) -> str:
    stem = Path(str(file_name or "スケジュール")).stem or "スケジュール"
    return f"{stem}_図.xlsx"


def gantt_columns(tasks, amount: int, unit: str) -> tuple[datetime, timedelta, int]:
    if amount < 1 or unit not in {"hour", "day"}:
        raise GanttRangeError("目盛は1以上の時間または日にしてください")
    origin = grid_origin(tasks)
    if origin is None:
        raise GanttRangeError("表示できる点がありません")
    step = step_size(amount, unit)
    latest = max(point.time for task in tasks for point in task.points)
    count = cell_index(latest, origin, step) + 1
    if count < 1 or count > MAX_CELLS:
        raise GanttRangeError("セルが多すぎます。目盛の間隔を広げてください")
    return origin, step, count


def header_text(start: datetime, step: timedelta, unit: str) -> str:
    if unit == "day":
        last = start + step - timedelta(days=1)
        if last.date() == start.date():
            return f"{start.month}/{start.day}"
        return f"{start.month}/{start.day}〜{last.month}/{last.day}"
    return f"{start.month}/{start.day} {start.hour:02d}:{start.minute:02d}"


def build_gantt_workbook(
    tasks,
    colors: dict[str, str],
    *,
    show_effort: bool,
    amount: int,
    unit: str,
    align: str | None = None,
    shade_effort: bool = False,
    start_align: str | None = None,
    middle_align: str | None = None,
    end_align: str | None = None,
    slot_names=None,
) -> Workbook:
    origin, step, count = gantt_columns(tasks, amount, unit)
    last_index = count - 1

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "図"
    thin = Border(
        left=Side(style="thin", color="D5DEE3"),
        right=Side(style="thin", color="D5DEE3"),
        top=Side(style="thin", color="D5DEE3"),
        bottom=Side(style="thin", color="D5DEE3"),
    )
    header_font = Font(name="Yu Gothic", size=8, color="5C6B73")
    name_font = Font(name="Yu Gothic", size=10, color="1B2428")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.cell(1, 1, "スロット").font = header_font
    sheet.cell(1, 1).alignment = center
    sheet.cell(1, 1).border = thin
    for index in range(last_index + 1):
        start = cell_start(index, origin, step)
        cell = sheet.cell(1, index + 2, header_text(start, step, unit))
        cell.font = header_font
        cell.alignment = center
        cell.border = thin
        sheet.column_dimensions[get_column_letter(index + 2)].width = 11 if unit == "day" else 6
    sheet.column_dimensions["A"].width = 18
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "B3" if show_effort else "B2"
    sheet.sheet_view.showGridLines = False

    row = 2
    if show_effort:
        sheet.cell(row, 1, "工数").font = name_font
        sheet.cell(row, 1).alignment = Alignment(vertical="center")
        sheet.cell(row, 1).border = thin
        spans = effort_spans(tasks)
        for index in range(last_index + 1):
            start = cell_start(index, origin, step)
            total = cell_effort(spans, start, start + step)
            cell = sheet.cell(row, index + 2, None if total is None else round(total, 1))
            cell.alignment = center
            cell.border = thin
            cell.font = Font(name="Yu Gothic", size=9, color="0E6B67")
        row += 1

    ordered = sorted(
        (task for task in tasks if task.points),
        key=lambda task: (task.slot, task.points[0].time, task.id),
    )
    labels = [slot_label(task.slot, slot_names) for task in ordered]
    if labels:
        widest = max(sum(2 if ord(char) > 127 else 1 for char in text) for text in labels)
        sheet.column_dimensions["A"].width = max(10, min(28, widest + 2))
    for task, label in zip(ordered, labels):
        sheet.cell(row, 1, label).font = name_font
        sheet.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row, 1).border = thin
        color = colors.get(task.id, "").lstrip("#")
        fills = (
            painted_cells(
                task,
                origin,
                step,
                align,
                color,
                shade_effort,
                start_align=start_align,
                middle_align=middle_align,
                end_align=end_align,
            )
            if color
            else {}
        )
        for index in range(last_index + 1):
            cell = sheet.cell(row, index + 2, None)
            cell.border = thin
            chosen = fills.get(index)
            if chosen:
                cell.fill = PatternFill("solid", fgColor=chosen.lstrip("#"))
        sheet.row_dimensions[row].height = 18
        row += 1
    return workbook


def write_gantt(path: str, tasks, colors: dict[str, str], **options) -> None:
    build_gantt_workbook(tasks, colors, **options).save(path)

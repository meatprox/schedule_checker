"""JS版の読み込みテストと同じ確認を、Python版のモデルに対して行う。"""

import unittest
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

from model import (
    Point,
    Segment,
    Task,
    assign_slots,
    bundled_sample_path,
    effort_spans,
    parse_workbook,
    sample_workbook,
    tasks_to_workbook,
)


def at(year, month, day, hour, minute):
    return datetime(year, month, day, hour, minute)


def workbook_from(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "予定"
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row_index, column, value)
    return workbook


def excel_serial(year, month, day, hour=0, minute=0):
    base = datetime(year, month, day, hour, minute) - datetime(1899, 12, 30)
    return base.total_seconds() / 86400


def make_point(hour, minute):
    return Point(id=f"{hour}-{minute}", time=at(2026, 10, 21, hour, minute), description="", row=0)


class ModelTests(unittest.TestCase):
    def test_datetime_cells(self):
        book = workbook_from(
            [
                ["タスク名", "日付", "時間", "説明"],
                ["朝会", datetime(2026, 10, 21), datetime(2026, 10, 21, 9, 0), "開始"],
                ["朝会", datetime(2026, 10, 21), datetime(2026, 10, 21, 9, 40), "終了"],
                ["移動", datetime(2026, 10, 21), datetime(2026, 10, 21, 10, 0), "出発"],
                ["移動", datetime(2026, 10, 21), datetime(2026, 10, 21, 10, 30), "到着"],
            ]
        )
        tasks = parse_workbook(book).tasks
        self.assertEqual([task.name for task in tasks], ["朝会", "移動"])
        self.assertEqual(tasks[0].points[0].time, at(2026, 10, 21, 9, 0))
        self.assertEqual(tasks[0].points[1].description, "終了")
        self.assertEqual(tasks[1].points[1].time, at(2026, 10, 21, 10, 30))

    def test_excel_serial(self):
        book = workbook_from(
            [
                ["タスク名", "日付", "時間", "説明"],
                ["朝会", excel_serial(2026, 10, 21), 9 / 24, "開始"],
            ]
        )
        sheet = book.active
        sheet["B2"].number_format = "yyyy/mm/dd"
        sheet["C2"].number_format = "hh:mm"
        task = parse_workbook(book).tasks[0]
        self.assertEqual(task.points[0].time, at(2026, 10, 21, 9, 0))

    def test_forward_fill(self):
        tasks = parse_workbook(
            workbook_from(
                [
                    ["朝会", "2026/10/21", "9:00", "開始"],
                    [None, None, "9:40", "A & B <終了>"],
                    ["移動", "10/21", "11時15分", "到着"],
                ]
            )
        ).tasks
        self.assertEqual(len(tasks), 2)
        self.assertEqual(len(tasks[0].points), 2)
        self.assertEqual(tasks[0].points[1].description, "A & B <終了>")
        self.assertEqual(tasks[1].points[0].time, at(2026, 10, 21, 11, 15))

    def test_blank_row_splits_same_name(self):
        tasks = parse_workbook(
            workbook_from(
                [
                    ["朝会", "2026年10月21日", "09:00", "開始"],
                    ["朝会", "2026年10月21日", "09:20", "終了"],
                    [None, None, None, None],
                    ["朝会", "2026年10月21日", "13:00", "午後の朝会"],
                ]
            )
        ).tasks
        self.assertEqual(len(tasks), 2)
        self.assertEqual(tasks[1].points[0].description, "午後の朝会")

    def test_title_row_reiwa_and_sort(self):
        tasks = parse_workbook(
            workbook_from(
                [
                    ["見出し"],
                    ["タスク名", "日付", "時間", "説明"],
                    ["打合せ", "令和8年10月21日", "午後1時30分", "開始"],
                    ["打合せ", "2026/10/21", "10:00", "先に書く遅い行"],
                ]
            )
        ).tasks
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0].points[0].time, at(2026, 10, 21, 10, 0))
        self.assertEqual(tasks[0].points[1].time, at(2026, 10, 21, 13, 30))

    def test_assign_slots(self):
        tasks = [
            Task(id="a", name="a", points=[make_point(9, 0), make_point(10, 0)]),
            Task(id="b", name="b", points=[make_point(9, 30), make_point(11, 0)]),
            Task(id="c", name="c", points=[make_point(10, 30), make_point(12, 0)]),
            Task(id="d", name="d", points=[make_point(9, 15), make_point(11, 30)]),
        ]
        assign_slots(tasks, 5)
        slots = {task.id: task.slot for task in tasks}
        self.assertEqual(slots, {"a": 0, "d": 1, "b": 2, "c": 0})

    def test_saved_sample_file(self):
        path = bundled_sample_path()
        self.assertTrue(path.exists(), path)
        self._assert_sample(parse_workbook(path))

    def test_embedded_sample_matches(self):
        self._assert_sample(parse_workbook(sample_workbook()))

    def _assert_sample(self, parsed):
        assign_slots(parsed.tasks, 5)
        names = {task.name: task for task in parsed.tasks}
        self.assertEqual(len(parsed.tasks), 10)
        morning = names["朝会"]
        self.assertEqual(len(morning.points), 3)
        self.assertEqual(morning.points[0].time.hour, 9)
        self.assertEqual(morning.points[1].description, "課題の洗い出し")
        self.assertNotEqual(names["設計レビュー"].slot, names["移動"].slot)
        self.assertEqual(len({names["実装"].slot, names["定例"].slot, names["相談"].slot}), 3)
        self.assertEqual(names["振り返り"].points[0].time.day, 22)
        self.assertNotIn("タスク名", names)
        spans = effort_spans(parsed.tasks)
        self.assertTrue(any(span.total == 2 for span in spans))
        self.assertTrue(any(span.total == 3 for span in spans))

    def test_export_round_trip(self):
        edited = [
            Task(
                id="t1",
                name="朝会",
                slot=2,
                points=[
                    Point("p1", at(2026, 10, 21, 11, 30), "開始", 1),
                    Point("p2", at(2026, 10, 21, 12, 0), "A & B <終了>", 2),
                ],
                segments=[Segment("dashed")],
            ),
            Task(
                id="t2",
                name="朝会",
                slot=4,
                points=[Point("p3", at(2026, 10, 22, 9, 5), "翌日", 4)],
            ),
        ]
        written = tasks_to_workbook(edited, "予定:テスト")
        self.assertEqual(written.sheetnames[0], "予定 テスト")
        self.assertEqual(written.active["F1"].value, "工数")
        self.assertEqual(written.active["F2"].value, 0.5)
        parsed = parse_workbook(written)
        assign_slots(parsed.tasks, 5)
        self.assertEqual(len(parsed.tasks), 2)
        self.assertEqual(parsed.tasks[0].points[0].time, at(2026, 10, 21, 11, 30))
        self.assertEqual(parsed.tasks[0].points[1].description, "A & B <終了>")
        self.assertEqual(parsed.tasks[0].segments[0].style, "dashed")
        self.assertEqual(parsed.tasks[0].slot, 2)
        self.assertTrue(parsed.tasks[0].slot_assigned)
        self.assertEqual(parsed.tasks[1].slot, 4)
        self.assertEqual(parsed.tasks[1].points[0].time, at(2026, 10, 22, 9, 5))
        self.assertEqual(
            parsed.tasks[1].points[0].time - parsed.tasks[0].points[0].time,
            timedelta(hours=21, minutes=35),
        )


if __name__ == "__main__":
    unittest.main()
